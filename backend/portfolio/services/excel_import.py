"""One-off import from the legacy spreadsheet (StockMarket_organizado.xlsx layout).

Tables are located by their header row, not by fixed coordinates, so small layout edits still import.
Positions that the spreadsheet derives from a ledger (crypto trades, metal purchases) are imported from
that ledger; every other position becomes a single opening BUY at its average cost.
"""

import datetime as dt
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation

import openpyxl
from django.db import transaction as db_tx
from django.utils import timezone

from portfolio.models import (
    Account,
    AccountType,
    AllocationTarget,
    Asset,
    AssetClass,
    CashFlow,
    EquityGrant,
    PriceSource,
    Snapshot,
    Transaction,
    TransactionKind,
)

ACCOUNT_TYPES = {
    "broker": AccountType.BROKERAGE,
    "roth ira": AccountType.ROTH_IRA,
    "401k": AccountType.K401,
    "crypto": AccountType.CRYPTO,
    "metales": AccountType.PHYSICAL,
    "efectivo": AccountType.BANK,
    "préstamo": AccountType.LOAN,
    "prestamo": AccountType.LOAN,
    "equity privada": AccountType.PRIVATE_EQUITY,
}

ASSET_CLASSES = {
    "acciones/etf": AssetClass.STOCK,
    "bonos": AssetClass.BOND,
    "crypto": AssetClass.CRYPTO,
    "metales": AssetClass.METAL,
    "efectivo": AssetClass.CASH,
    "deuda": AssetClass.DEBT,
    "equity privada": AssetClass.PRIVATE_EQUITY,
}

COINGECKO_IDS = {"BTC": "bitcoin", "ETH": "ethereum", "XRP": "ripple", "SOL": "solana"}

# Spreadsheet metal names -> (symbol, Yahoo futures ticker, display name)
METALS = {
    "oro": ("XAU", "GC=F", "Oro (oz)"),
    "plata": ("XAG", "SI=F", "Plata (oz)"),
    "platino": ("XPT", "PL=F", "Platino (oz)"),
}


@dataclass
class ImportReport:
    accounts: int = 0
    assets: int = 0
    transactions: int = 0
    cash_flows: int = 0
    targets: int = 0
    equity_grants: int = 0
    snapshots: int = 0
    warnings: list[str] = field(default_factory=list)


def _dec(v) -> Decimal | None:
    if v is None or isinstance(v, str) and (v.startswith("=") or not v.strip()):
        return None
    try:
        return Decimal(str(v))
    except InvalidOperation:
        return None


def _date(v) -> dt.date | None:
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    if isinstance(v, str):
        for fmt in ("%m/%d/%Y", "%Y-%m-%d"):
            try:
                return dt.datetime.strptime(v.strip(), fmt).date()
            except ValueError:
                pass
    return None


def _table(ws, first_header: str):
    """Yield dict rows for the table whose header row starts with `first_header` in column A."""
    rows = list(ws.iter_rows(values_only=True))
    for i, row in enumerate(rows):
        if row and row[0] == first_header:
            headers = [str(h).strip() if h is not None else None for h in row]
            for data in rows[i + 1 :]:
                if all(c is None for c in data[:4]):
                    return
                if isinstance(data[0], str) and data[0].strip().upper() in {"TOTAL", "TOTAL GENERAL"}:
                    return
                yield {h: v for h, v in zip(headers, data) if h}
            return


class LegacyImporter:
    def __init__(self, path, user, fallback_date: dt.date | None = None):
        self.wb = openpyxl.load_workbook(path, data_only=False)
        self.user = user
        self.fallback_date = fallback_date or timezone.localdate()
        self.report = ImportReport()
        self.accounts: dict[str, Account] = {}
        self.assets: dict[str, Asset] = {}

    def run(self, reset: bool = False) -> ImportReport:
        with db_tx.atomic():
            if reset:
                self._reset()
            self._accounts()
            self._positions()
            self._crypto_trades()
            self._crypto_deposits()
            self._metals()
            self._equity()
            self._history()
        return self.report

    def _reset(self):
        Transaction.objects.filter(account__owner=self.user).delete()
        Account.objects.filter(owner=self.user).delete()
        Asset.objects.filter(owner=self.user).delete()
        Snapshot.objects.filter(owner=self.user).delete()

    # -- helpers ----------------------------------------------------------------------------------

    def _asset(self, symbol: str, **defaults) -> Asset:
        if symbol in self.assets:
            return self.assets[symbol]
        asset, created = Asset.objects.get_or_create(owner=self.user, symbol=symbol, defaults=defaults)
        self.report.assets += created
        self.assets[symbol] = asset
        return asset

    def _buy(self, account, asset, qty, price, date=None, **kw):
        Transaction.objects.create(
            account=account,
            asset=asset,
            kind=TransactionKind.BUY,
            date=date or self.fallback_date,
            quantity=qty,
            price=price,
            **kw,
        )
        self.report.transactions += 1

    # -- sheets -----------------------------------------------------------------------------------

    def _accounts(self):
        for row in _table(self.wb["Resumen"], "Cuenta"):
            name = row["Cuenta"]
            acc_type = ACCOUNT_TYPES.get(str(row.get("Tipo") or "").lower(), AccountType.OTHER)
            acc, created = Account.objects.get_or_create(
                owner=self.user,
                name=name,
                defaults={
                    "institution": row.get("Institución") or "",
                    "account_type": acc_type,
                    "is_liquid": str(row.get("Accesible") or "").lower().startswith("s"),
                },
            )
            self.report.accounts += created
            self.accounts[name] = acc

    def _positions(self):
        for row in _table(self.wb["Posiciones"], "Cuenta"):
            account = self.accounts.get(row["Cuenta"])
            asset_class = ASSET_CLASSES.get(str(row.get("Clase") or "").lower(), AssetClass.OTHER)
            symbol = str(row.get("Símbolo") or "").strip()
            if not account or not symbol:
                continue
            if asset_class in (AssetClass.CRYPTO, AssetClass.METAL) and _dec(row.get("Cantidad")) is None:
                continue  # derived from the trades/metals ledgers, imported separately
            if asset_class == AssetClass.PRIVATE_EQUITY:
                continue  # imported from the Equity sheet

            manual = _dec(row.get("Precio manual"))
            name = row.get("Descripción") or symbol
            if asset_class == AssetClass.CASH:
                symbol, name, source, manual = "USD", "Dólar", PriceSource.MANUAL, Decimal(1)
            elif asset_class == AssetClass.DEBT:
                symbol = f"LOAN-{account.name.upper().replace(' ', '-')}"
                source, manual = PriceSource.MANUAL, Decimal(1)
            else:
                # Rows without a live-price formula carry a manual price (e.g. 401k funds with no public ticker).
                source = PriceSource.MANUAL if manual is not None else PriceSource.YAHOO

            asset = self._asset(
                symbol,
                name=name,
                asset_class=asset_class,
                price_source=source,
                manual_price=manual if source == PriceSource.MANUAL else None,
            )
            qty, cost = _dec(row.get("Cantidad")), _dec(row.get("Costo prom."))
            if qty:
                self._buy(account, asset, qty, cost or Decimal(0), notes="Saldo inicial (importado de Excel)")

            target = _dec(row.get("Objetivo %"))
            if target is not None and asset_class != AssetClass.CASH:
                AllocationTarget.objects.update_or_create(
                    account=account, asset=asset, defaults={"target_pct": target * 100}
                )
                self.report.targets += 1

    def _crypto_trades(self):
        account = self.accounts.get("Crypto")
        if not account or "Crypto-Trades" not in self.wb.sheetnames:
            return
        for row in _table(self.wb["Crypto-Trades"], "Fecha"):
            symbol = str(row.get("Activo") or "").strip().upper()
            paid, qty = _dec(row.get("Pagado (USD)")), _dec(row.get("Cantidad"))
            if not symbol or not paid or not qty:
                continue
            asset = self._asset(
                symbol,
                name=symbol,
                asset_class=AssetClass.CRYPTO,
                price_source=PriceSource.COINGECKO,
                source_id=COINGECKO_IDS.get(symbol, symbol.lower()),
            )
            date = _date(row.get("Fecha"))
            notes = [] if date else ["Fecha desconocida en Excel"]
            if row.get("Exchange"):
                notes.append(f"Exchange: {row['Exchange']}")
            self._buy(
                account, asset, qty, paid / qty, date=date,
                external_id=str(row.get("Order ID") or ""), notes="; ".join(notes),
            )
            if not date:
                self.report.warnings.append(f"Trade {symbol} {qty} sin fecha → {self.fallback_date}")

    def _crypto_deposits(self):
        account = self.accounts.get("Crypto")
        if not account or "Crypto" not in self.wb.sheetnames:
            return
        for row in _table(self.wb["Crypto"], "Exchange / concepto"):
            amount = _dec(row.get("Depósito"))
            if not amount:
                continue
            CashFlow.objects.create(
                account=account,
                date=_date(row.get("Fecha")) or self.fallback_date,
                amount=amount,
                notes=str(row.get("Exchange / concepto") or ""),
            )
            self.report.cash_flows += 1

    def _metals(self):
        account = self.accounts.get("Metales")
        if not account or "Metales" not in self.wb.sheetnames:
            return
        for row in _table(self.wb["Metales"], "Fecha"):
            metal = METALS.get(str(row.get("Metal") or "").strip().lower())
            qty, total = _dec(row.get("Cantidad (oz)")), _dec(row.get("Costo total"))
            if not metal or not qty or total is None:
                continue
            symbol, yahoo, name = metal
            asset = self._asset(
                symbol, name=name, asset_class=AssetClass.METAL, price_source=PriceSource.YAHOO, source_id=yahoo
            )
            self._buy(account, asset, qty, total / qty, date=_date(row.get("Fecha")), notes=row.get("Nota") or "")

    def _equity(self):
        if "Equity" not in self.wb.sheetnames:
            return
        for row in _table(self.wb["Equity"], "Empresa"):
            qty, strike, fmv = _dec(row.get("Opciones")), _dec(row.get("Strike")), _dec(row.get("FMV por acción"))
            if not qty or strike is None or fmv is None:
                continue
            company = row["Empresa"]
            acc = self.accounts.get(f"{company} Equity")
            if acc is None:
                acc, _ = Account.objects.get_or_create(
                    owner=self.user, name=f"{company} Equity",
                    defaults={"account_type": AccountType.PRIVATE_EQUITY, "is_liquid": False, "institution": company},
                )
            EquityGrant.objects.create(
                account=acc,
                company=company,
                quantity=qty,
                strike=strike,
                fmv=fmv,
                exercised=str(row.get("¿Ejercidas?") or "").lower().startswith("s"),
            )
            self.report.equity_grants += 1

    def _history(self):
        if "Historial" not in self.wb.sheetnames:
            return
        ws = self.wb["Historial"]
        for row in _table(ws, "Año"):
            year, value = row.get("Año"), _dec(row.get("Patrimonio"))
            if value is None or not str(year).isdigit():
                continue
            Snapshot.objects.update_or_create(
                owner=self.user, date=dt.date(int(year), 12, 31),
                defaults={"net_worth": value, "notes": row.get("Nota") or "Fin de año (importado)"},
            )
            self.report.snapshots += 1
        for row in _table(ws, "Fecha"):
            date, value = _date(row.get("Fecha")), _dec(row.get("Patrimonio total"))
            if not date or value is None:
                continue
            Snapshot.objects.update_or_create(
                owner=self.user, date=date,
                defaults={
                    "net_worth": value,
                    "net_worth_ex_private": _dec(row.get("Sin equity")),
                    "liquid": _dec(row.get("Líquido")),
                    "cost": _dec(row.get("Costo")),
                    "unrealized": _dec(row.get("Ganancia")),
                    "notes": "Foto (importado)",
                },
            )
            self.report.snapshots += 1
